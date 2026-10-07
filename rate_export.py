"""
rate_export.py - 汇率整合导出
================================

从 BOCHK（中银香港）网站获取港币电汇牌价，参照 TFISF 汇率表
的公式计算自定义汇率（Bid Mark down / Ask Mark up），
并从 HKEx 获取交易所印花税率，按模板格式生成两个 xlsx 文件：

  1. 自定义汇率yymmdd.xlsx — 来源: BOCHK 港币电汇牌价（按公式计算）
  2. 交易所汇率yymmdd.xlsx — 来源: 仅 HKEx 印花税率

自定义汇率计算公式（参照 TFISF Excel）：
  F = Bid  = BOCHK 客户卖出价 (银行买入价)
  H = Ask  = BOCHK 客户买入价 (银行卖出价)
  L = Bid (Mark down) = F * markdown_factor  → 模板 Buy
  N = Ask (Mark up)    = H * markup_factor    → 模板 Sell
  USD/CNY 等交叉汇率来自 BOCHK 美元电汇牌价页面

模板格式 (Ledger_Spot Rate Import):
  Sheet 名: 各国CNY、HKD汇率
  Row 0: 日期Date | <excel 日期序列号>
  Row 1: 源币种Currency | 目标币种  To_Currency |
         源币种单位 Currency Unit |
         目标币种等值(买入)  Equivalent(Buy) |
         目标币种等值(卖出)  Equivalent(Sell)
  Row 2+: 数据行 (CNY/HKD, USD/HKD, USD/CNY, JPY/HKD, ...)

仅在企业微信 webhook 上发送异常通知；正常情况下不发 webhook。
"""

from __future__ import annotations

import datetime as dt
import logging
import os
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from sc_infra import send_error
from web_rates import (
    fetch_bochk_fxrates,
    fetch_bochk_hkdrates,
    fetch_bochk_usdrates,
    fetch_hkex_stampfx,
)
from web_shot import capture_rate_pages

logger = logging.getLogger(__name__)


def _env(key: str) -> str:
    return os.getenv(key, "").strip().strip('"').strip("'")


# ---------------------------------------------------------------------------
# 汇率对白名单配置
# ---------------------------------------------------------------------------

# 保持 env 中书写顺序：既作白名单，也作导出行的排列顺序
RATE_PAIRS: list[str] = [
    s.strip().upper()
    for s in _env("RATE_PAIRS").split(",")
    if s.strip()
]

# RATE_PAIRS 未配置时的默认排列顺序（原硬编码优先级）
DEFAULT_RATE_PAIR_ORDER: tuple[str, ...] = (
    "CNY/HKD",
    "USD/HKD",
    "USD/CNY",
)


def _is_pair_allowed(from_ccy: str, to_ccy: str) -> bool:
    """检查汇率对是否在 RATE_PAIRS 白名单中。

    RATE_PAIRS 为空时表示不限制，全部允许。
    RATE_PAIRS 的书写顺序同时决定导出行的排列顺序。

    Args:
        from_ccy: 源币种代码
        to_ccy: 目标币种代码

    Returns:
        True 表示允许写入
    """
    if not RATE_PAIRS:
        return True
    return f"{from_ccy}/{to_ccy}".upper() in RATE_PAIRS


# ---------------------------------------------------------------------------
# 模板常量
# ---------------------------------------------------------------------------

SHEET_NAME = "各国CNY、HKD汇率"

TEMPLATE_HEADERS: list[str] = [
    "源币种Currency",
    "目标币种  To_Currency",
    "源币种单位 Currency Unit",
    "目标币种等值(买入)  Equivalent(Buy)",
    "目标币种等值(卖出)  Equivalent(Sell)",
]

# BOCHK 港币牌价 货币中文名 -> ISO 币种代码
BOCHK_CURRENCY_MAP: dict[str, str] = {
    "人民幣(在岸)": "CNY",
    "人民幣(離岸)": "CNH",
    "美元": "USD",
    "英鎊": "GBP",
    "日圓": "JPY",
    "澳元": "AUD",
    "紐元": "NZD",
    "加元": "CAD",
    "歐羅": "EUR",
    "瑞士法郎": "CHF",
    "丹麥克郎": "DKK",
    "挪威克郎": "NOK",
    "瑞典克郎": "SEK",
    "新加坡元": "SGD",
    "泰國銖": "THB",
    "文萊元": "BND",
    "南非蘭特": "ZAR",
}

# BOCHK 现钞牌价 货币中文名 -> ISO 币种代码
# 现钞牌价页比电汇牌价页多出部分币种（如韓國圜 KRW、新台幣 TWD 等），
# 电汇牌价取不到的币对从这里补充。
BOCHK_FX_CURRENCY_MAP: dict[str, str] = {
    "人民幣": "CNY",
    "澳元": "AUD",
    "加拿大元": "CAD",
    "瑞士法郎": "CHF",
    "歐羅": "EUR",
    "英鎊": "GBP",
    "日圓": "JPY",
    "紐西蘭元": "NZD",
    "新加坡元": "SGD",
    "泰國銖": "THB",
    "美元": "USD",
    "印尼盾": "IDR",
    "印度盧比": "INR",
    "韓國圜": "KRW",
    "澳門元": "MOP",
    "菲律賓彼索": "PHP",
    "俄羅斯盧布": "RUB",
    "新台幣": "TWD",
    "文萊元": "BND",
    "南非蘭特": "ZAR",
}

# BOCHK 美元牌价 货币对中文名 -> (FROM_ISO, TO_ISO)
# 格式为 "源币/目标币"，如 "美元/人民幣" -> ("USD", "CNY")
BOCHK_USD_PAIR_MAP: dict[str, tuple[str, str]] = {
    "美元/人民幣": ("USD", "CNY"),
    "美元/港元": ("USD", "HKD"),
    "美元/加元": ("USD", "CAD"),
    "美元/瑞士法郎": ("USD", "CHF"),
    "美元/丹麥克郎": ("USD", "DKK"),
    "美元/日元": ("USD", "JPY"),
    "美元/挪威克郎": ("USD", "NOK"),
    "美元/瑞典克郎": ("USD", "SEK"),
    "美元/新加坡元": ("USD", "SGD"),
    "美元/泰國銖": ("USD", "THB"),
    "美元/文萊元": ("USD", "BND"),
    "美元/南非蘭特": ("USD", "ZAR"),
    "澳元/美元": ("AUD", "USD"),
    "歐羅/美元": ("EUR", "USD"),
    "英鎊/美元": ("GBP", "USD"),
    "紐元/美元": ("NZD", "USD"),
}

# Markdown / Markup 因子配置（参照 TFISF Excel 2026 sheet 公式）
# key = "FROM/TO", value = (markdown_factor, markup_factor)
#   L (Bid Mark down) = F * markdown_factor  → 模板 Buy
#   N (Ask Mark up)    = H * markup_factor    → 模板 Sell
MARKDOWN_FACTORS: dict[str, tuple[float, float]] = {
    "CNY/HKD": (0.98, 1.02),
    "USD/HKD": (0.998, 1.002),
    "USD/CNY": (0.98, 1.02),
}
# 默认因子（未配置的币对）：1.0 表示不加价差，
# Buy/Sell 直接等于 BOCHK 原始牌价
DEFAULT_MARKDOWN_FACTOR: tuple[float, float] = (1.0, 1.0)

# HKEx 货币英文名 -> ISO 币种代码
HKEX_CURRENCY_MAP: dict[str, str] = {
    "Renminbi": "CNY",
    "U.S. dollars": "USD",
}


# ---------------------------------------------------------------------------
# 数据收集
# ---------------------------------------------------------------------------

# 匹配资料更新时间中的日期部分，如 "2026/09/29 13:42:59"
_UPDATE_DATE_RE = re.compile(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})")


def _fetch_safe(label: str, errors: list[str], func, *args) -> dict:
    """安全调用抓取函数，将网络/解析异常记录为错误。

    任何「页面打不开、请求失败、解析失败」的异常都会转成一条
    错误信息，不会中断整个流程。

    Args:
        label: 数据源名称，用于错误信息前缀
        errors: 错误信息收集列表
        func: 抓取函数
        *args: 传给抓取函数的位置参数

    Returns:
        抓取结果的字典，失败时为空字典
    """
    try:
        return func(*args) or {}
    except Exception as exc:
        msg = f"{label}: 页面无法打开或抓取失败 ({type(exc).__name__}: {exc})"
        errors.append(msg)
        logger.error("[RE] %s", msg)
        return {}


# 香港公众假期（用于判断 HKEx 是否开市）。加载失败则回退到仅按工作日判断，
# 以免依赖缺失导致整个流程崩溃。
try:
    import holidays as _holidays_lib
    _HK_HOLIDAYS = _holidays_lib.country_holidays(
        "HK", years=range(2020, 2041),
    )
except Exception:  # pragma: no cover - 依赖缺失时仅按工作日判断
    _HK_HOLIDAYS = None
    logger.warning(
        "未能加载 holidays 库，将仅按周一至周五判断 HKEx 交易日（不含公众假期）"
    )


def _is_hk_trading_day(date_str: str) -> bool:
    """判断日期是否为 HKEx 交易日（周一至周五 且非香港公众假期）。

    HKEx 在周末及香港公众假期不发布印花税率，此时取不到当日数据属于
    正常现象，不应当作错误。
    中银香港(BOCHK)在假期仍可能提供参考汇率，因此「非交易日」也用于放宽
    对 BOCHK 的当日校验。

    Args:
        date_str: YYYYMMDD 格式日期

    Returns:
        True 表示 HKEx 应当开市并发布当日汇率
    """
    try:
        d = dt.datetime.strptime(date_str, "%Y%m%d").date()
    except ValueError:
        return True
    if d.weekday() >= 5:
        return False
    if _HK_HOLIDAYS is not None and d in _HK_HOLIDAYS:
        return False
    return True


def _check_update_date(
    label: str,
    date_str: str,
    update_time: str,
    errors: list[str],
) -> None:
    """校验数据的资料更新时间是否为目标日期，不是当日则记为错误。

    目标日期本身非交易日（周末/香港公众假期）时，HKEx/BOCHK 不会发布
    当日汇率，此时只记录日志，不算错误。

    Args:
        label: 数据源名称
        date_str: 目标日期 YYYYMMDD
        update_time: 页面上的资料更新时间字符串
        errors: 错误信息收集列表
    """
    if not update_time:
        errors.append(f"{label}: 未获取到资料更新时间，无法确认是否为当日汇率")
        return

    m = _UPDATE_DATE_RE.search(update_time)
    if not m:
        errors.append(f"{label}: 资料更新时间格式无法解析: {update_time}")
        return

    upd = f"{int(m.group(1)):04d}{int(m.group(2)):02d}{int(m.group(3)):02d}"
    if upd == date_str:
        return

    msg = (
        f"{label}: 汇率不是当日的 "
        f"(资料更新时间 {update_time}, 期望日期 {date_str})"
    )
    if _is_hk_trading_day(date_str):
        errors.append(msg)
        logger.warning("[RE] %s", msg)
    else:
        logger.info("[RE] %s (目标日期非交易日, 仅记录)", msg)


def _collect_bochk_rates(date_str: str, errors: list[str]) -> list[dict]:
    """从 BOCHK 网站获取港币电汇牌价。

    Args:
        date_str: 目标日期 YYYYMMDD，用于校验是否为当日汇率
        errors: 错误信息收集列表

    Returns:
        汇率字典列表，每项含 currency / sell / buy
    """
    data = _fetch_safe("BOCHK 港币电汇牌价", errors, fetch_bochk_hkdrates)
    rates = data.get("rates", [])
    if not rates:
        errors.append("BOCHK: 未获取到港币电汇牌价数据")
    _check_update_date(
        "BOCHK 港币电汇牌价", date_str, data.get("update_time", ""), errors,
    )
    return rates


def _collect_bochk_usd_rates(date_str: str, errors: list[str]) -> list[dict]:
    """从 BOCHK 网站获取美元电汇牌价。

    美元牌价页面提供各币种兑美元的汇率，
    货币名格式为 "美元/人民幣" 等。

    Args:
        date_str: 目标日期 YYYYMMDD，用于校验是否为当日汇率
        errors: 错误信息收集列表

    Returns:
        汇率字典列表，每项含 currency / sell / buy
    """
    data = _fetch_safe("BOCHK 美元电汇牌价", errors, fetch_bochk_usdrates)
    rates = data.get("rates", [])
    if not rates:
        errors.append("BOCHK-USD: 未获取到美元电汇牌价数据")
    _check_update_date(
        "BOCHK 美元电汇牌价", date_str, data.get("update_time", ""), errors,
    )
    return rates


def _collect_bochk_fx_rates(date_str: str, errors: list[str]) -> list[dict]:
    """从 BOCHK 网站获取港币现钞牌价。

    现钞牌价页包含电汇牌价页没有的币种（如韓國圜 KRW、新台幣 TWD），
    用于补充电汇牌价取不到的 XXX/HKD 币对。

    Args:
        date_str: 目标日期 YYYYMMDD，用于校验是否为当日汇率
        errors: 错误信息收集列表

    Returns:
        汇率字典列表，每项含 currency / sell / buy
    """
    data = _fetch_safe("BOCHK 港币现钞牌价", errors, fetch_bochk_fxrates)
    rates = data.get("rates", [])
    if not rates:
        errors.append("BOCHK-FX: 未获取到港币现钞牌价数据")
    _check_update_date(
        "BOCHK 港币现钞牌价", date_str, data.get("update_time", ""), errors,
    )
    return rates


def _collect_hkex_rates(date_str: str, errors: list[str]) -> dict:
    """从 HKEx 网站获取印花税率汇率。

    HKEx 仅在交易日（周一至五且非香港公众假期）发布印花税率，
    且只取目标日期当天的数据，绝不回退到上一交易日。
    取不到数据时：交易日记为错误，非交易日属正常现象不记错。

    Args:
        date_str: YYYYMMDD 格式日期
        errors: 错误信息收集列表

    Returns:
        包含 date / rates / xls_url 的字典，可能为空
    """
    data = _fetch_safe("HKEx 印花税率", errors, fetch_hkex_stampfx, date_str)
    if not data:
        # HKEx 在周末及香港公众假期不发布印花税率，属正常现象，不当作错误。
        if _is_hk_trading_day(date_str):
            errors.append(f"HKEx: 未获取到印花税率数据 (日期 {date_str})")
        else:
            logger.info(
                "[RE] HKEx 在 %s 非交易日(周末/香港公众假期), "
                "无印花税率数据, 跳过",
                date_str,
            )
        return {}

    # fetch_hkex_stampfx 只返回目标日期当天的数据，取不到就是空，
    # 不存在「抓到别的日期」的情况，无需再校验日期是否一致。
    if not data.get("rates"):
        errors.append(f"HKEx: 印花税率数据为空 (日期 {date_str})")
    return data


# ---------------------------------------------------------------------------
# 数据转换
# ---------------------------------------------------------------------------

def _bochk_to_raw_dict(rates: list[dict]) -> dict[str, dict[str, float]]:
    """将 BOCHK 港币牌价解析为 {币种代码: {"bid": F, "ask": H}} 字典。

    BOCHK 的 sell 是客户卖出价(银行买入价)，对应 Excel 的 Bid(F)。
    BOCHK 的 buy 是客户买入价(银行卖出价)，对应 Excel 的 Ask(H)。

    Args:
        rates: BOCHK 汇率列表，每项含 currency / sell / buy

    Returns:
        {币种代码: {"bid": float, "ask": float}} 字典
    """
    result: dict[str, dict[str, float]] = {}
    for r in rates:
        ccy_name = r.get("currency", "")
        ccy_code = BOCHK_CURRENCY_MAP.get(ccy_name)
        if not ccy_code:
            logger.warning("[RE] BOCHK 未知货币: %s, 跳过", ccy_name)
            continue
        bid = _to_float(r.get("sell"))
        ask = _to_float(r.get("buy"))
        result[ccy_code] = {"bid": bid, "ask": ask}
    return result


def _bochk_usd_to_raw_dict(
    rates: list[dict],
) -> dict[str, dict[str, float]]:
    """将 BOCHK 美元牌价解析为 {"FROM/TO": {"bid": F, "ask": H}} 字典。

    美元牌价的 currency 格式为 "美元/人民幣" 等，
    通过 BOCHK_USD_PAIR_MAP 映射为 ISO 币种对。

    Args:
        rates: BOCHK 美元牌价列表，每项含 currency / sell / buy

    Returns:
        {"USD/CNY": {"bid": float, "ask": float}} 字典
    """
    result: dict[str, dict[str, float]] = {}
    for r in rates:
        ccy_name = r.get("currency", "")
        pair = BOCHK_USD_PAIR_MAP.get(ccy_name)
        if not pair:
            logger.warning("[RE] BOCHK-USD 未知货币对: %s, 跳过", ccy_name)
            continue
        pair_key = f"{pair[0]}/{pair[1]}"
        bid = _to_float(r.get("sell"))
        ask = _to_float(r.get("buy"))
        result[pair_key] = {"bid": bid, "ask": ask}
    return result


def _bochk_fx_to_raw_dict(
    rates: list[dict],
) -> dict[str, dict[str, float]]:
    """将 BOCHK 现钞牌价解析为 {币种代码: {"bid": F, "ask": H}} 字典。

    Args:
        rates: BOCHK 现钞牌价列表，每项含 currency / sell / buy

    Returns:
        {币种代码: {"bid": float, "ask": float}} 字典
    """
    result: dict[str, dict[str, float]] = {}
    for r in rates:
        ccy_name = r.get("currency", "")
        ccy_code = BOCHK_FX_CURRENCY_MAP.get(ccy_name)
        if not ccy_code:
            logger.warning("[RE] BOCHK-FX 未知货币: %s, 跳过", ccy_name)
            continue
        bid = _to_float(r.get("sell"))
        ask = _to_float(r.get("buy"))
        result[ccy_code] = {"bid": bid, "ask": ask}
    return result


def _apply_markdown(
    pair_key: str,
    bid: float,
    ask: float,
) -> tuple[float, float]:
    """参照 TFISF Excel 公式计算 Buy (Mark down) 和 Sell (Mark up)。

    L = Bid * markdown_factor  → Buy
    N = Ask * markup_factor    → Sell

    Args:
        pair_key: 币种对，如 "CNY/HKD"
        bid: BOCHK 客户卖出价 (银行买入价)
        ask: BOCHK 客户买入价 (银行卖出价)

    Returns:
        (buy, sell) 元组
    """
    md_factor, mu_factor = MARKDOWN_FACTORS.get(
        pair_key, DEFAULT_MARKDOWN_FACTOR,
    )
    return bid * md_factor, ask * mu_factor


def _compute_custom_rows(
    bochk_hkd: dict[str, dict[str, float]],
    bochk_usd: dict[str, dict[str, float]],
    bochk_fx: dict[str, dict[str, float]] | None = None,
) -> list[dict]:
    """根据 BOCHK 原始汇率，参照 TFISF Excel 公式计算自定义汇率。

    计算规则（参照 TFISF Excel 2026 sheet）：
      F = Bid（BOCHK 客户卖出价 = 银行买入价）
      H = Ask（BOCHK 客户买入价 = 银行卖出价）
      L = Bid (Mark down) = F * markdown_factor  → 模板 Buy
      N = Ask (Mark up)    = H * markup_factor    → 模板 Sell

    数据来源:
      - 各币种/HKD: 来自 BOCHK 港币电汇牌价
      - USD/CNY 等: 来自 BOCHK 美元电汇牌价
      USD/HKD 不会从美元牌价重复添加（港币牌价已有）
      - 电汇牌价没有的币种（如 KRW）: 来自 BOCHK 港币现钞牌价，
        仅在 RATE_PAIRS 显式配置了该币对时补充

    Args:
        bochk_hkd: {币种代码: {"bid": F, "ask": H}} 港币电汇牌价字典
        bochk_usd: {"FROM/TO": {"bid": F, "ask": H}} 美元电汇牌价字典
        bochk_fx: {币种代码: {"bid": F, "ask": H}} 港币现钞牌价字典，可选

    Returns:
        模板行字典列表，含 from_ccy / to_ccy / unit / buy / sell
    """
    rows: list[dict] = []
    seen: set[str] = set()

    # 1. 各币种兑 HKD 的直接汇率（港币牌价）
    for ccy_code, val in bochk_hkd.items():
        pair_key = f"{ccy_code}/HKD"
        seen.add(pair_key)
        buy, sell = _apply_markdown(pair_key, val["bid"], val["ask"])
        rows.append({
            "from_ccy": ccy_code,
            "to_ccy": "HKD",
            "unit": 1,
            "buy": buy,
            "sell": sell,
        })

    # 2. 美元牌价中的交叉汇率对（如 USD/CNY）
    #    跳过 USD/HKD 等已在港币牌价中存在的对
    for pair_key, val in bochk_usd.items():
        if pair_key in seen:
            continue
        seen.add(pair_key)
        from_ccy, to_ccy = pair_key.split("/", 1)
        buy, sell = _apply_markdown(pair_key, val["bid"], val["ask"])
        rows.append({
            "from_ccy": from_ccy,
            "to_ccy": to_ccy,
            "unit": 1,
            "buy": buy,
            "sell": sell,
        })

    # 3. 现钞牌价补充：电汇牌价没有的币种（如 KRW）
    #    仅在 RATE_PAIRS 显式配置了该币对时补充，避免未配置白名单时
    #    把现钞页的多余币种（IDR/INR/MOP/PHP/RUB/TWD 等）也导出。
    if bochk_fx and RATE_PAIRS:
        for ccy_code, val in bochk_fx.items():
            pair_key = f"{ccy_code}/HKD"
            if pair_key in seen or not _is_pair_allowed(ccy_code, "HKD"):
                continue
            seen.add(pair_key)
            buy, sell = _apply_markdown(pair_key, val["bid"], val["ask"])
            rows.append({
                "from_ccy": ccy_code,
                "to_ccy": "HKD",
                "unit": 1,
                "buy": buy,
                "sell": sell,
            })
            logger.info("[RE] %s 由 BOCHK 现钞牌价补充", pair_key)

    return rows


def _hkex_to_rows(data: dict) -> list[dict]:
    """将 HKEx 印花税率数据转换为模板行格式。

    HKEx 只提供 hkd_rate (中间价)，Buy 和 Sell 均设为该值。

    Args:
        data: HKEx 返回的数据字典

    Returns:
        模板行字典列表
    """
    rows: list[dict] = []
    for r in data.get("rates", []):
        name = r.get("currency", "")
        ccy_code = HKEX_CURRENCY_MAP.get(name)
        if not ccy_code:
            logger.warning("[RE] HKEx 未知货币: %s, 跳过", name)
            continue
        rate = r.get("hkd_rate", 0.0)
        rows.append({
            "from_ccy": ccy_code,
            "to_ccy": "HKD",
            "unit": r.get("unit", 1),
            "buy": rate,
            "sell": rate,
        })
    return rows


def _to_float(val) -> float:
    """安全转换为 float，空值返回 0.0。"""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return 0.0
    return float(s)


def _filter_rows(rows: list[dict]) -> list[dict]:
    """过滤掉不在 RATE_PAIRS 白名单中的汇率行。

    Args:
        rows: 汇率行字典列表

    Returns:
        过滤后的行列表
    """
    if not RATE_PAIRS:
        return rows
    return [
        r for r in rows
        if _is_pair_allowed(r["from_ccy"], r["to_ccy"])
    ]


def _missing_pairs(rows: list[dict]) -> list[str]:
    """返回 RATE_PAIRS 中配置了、但实际没取到数据的币对。

    RATE_PAIRS 只是过滤白名单，抓不到的币对会被静默丢掉；
    本函数用于把这种情况显式暴露出来，避免再出现「配了却没输出」。

    Args:
        rows: 汇率行字典列表

    Returns:
        缺失的币对字符串列表，RATE_PAIRS 未配置时恒为空
    """
    if not RATE_PAIRS:
        return []
    got = {f"{r['from_ccy']}/{r['to_ccy']}" for r in rows}
    return [p for p in RATE_PAIRS if p not in got]


# ---------------------------------------------------------------------------
# 数据排序
# ---------------------------------------------------------------------------

def _sort_rows(rows: list[dict]) -> list[dict]:
    """按 RATE_PAIRS 配置顺序排列汇率行。

    配置了 RATE_PAIRS 时，严格按 env 中书写的先后顺序排列；
    未配置时回退到默认顺序 CNY/HKD, USD/HKD, USD/CNY。
    两种情况中，未列出的币对都排在后面并按 from_ccy/to_ccy 字母序。

    Args:
        rows: 汇率行列表

    Returns:
        排序后的行列表
    """
    order = RATE_PAIRS or list(DEFAULT_RATE_PAIR_ORDER)
    rank = {pair: idx for idx, pair in enumerate(order)}

    def sort_key(row: dict) -> tuple:
        key = f"{row['from_ccy']}/{row['to_ccy']}"
        pri = rank.get(key, len(order))
        return (pri, row["from_ccy"], row["to_ccy"])

    return sorted(rows, key=sort_key)


# ---------------------------------------------------------------------------
# Excel 生成
# ---------------------------------------------------------------------------

def _write_xlsx(
    rows: list[dict],
    output_path: Path,
    date_val: dt.date,
) -> None:
    """按模板格式生成 xlsx 文件。

    Args:
        rows: 汇率行字典列表
        output_path: 输出文件路径
        date_val: 汇率日期
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    # Row 0: 日期Date | 日期值
    ws.cell(row=1, column=1, value="日期Date")
    date_cell = ws.cell(row=1, column=2, value=date_val)
    date_cell.number_format = "yyyy/m/d"

    # Row 1: 表头
    for col_idx, header in enumerate(TEMPLATE_HEADERS, start=1):
        ws.cell(row=2, column=col_idx, value=header)

    # Row 2+: 数据
    for row_idx, row_data in enumerate(rows, start=3):
        ws.cell(row=row_idx, column=1, value=row_data["from_ccy"])
        ws.cell(row=row_idx, column=2, value=row_data["to_ccy"])
        ws.cell(row=row_idx, column=3, value=row_data["unit"])
        buy_cell = ws.cell(row=row_idx, column=4, value=row_data["buy"])
        buy_cell.number_format = "0.0000000"
        sell_cell = ws.cell(row=row_idx, column=5, value=row_data["sell"])
        sell_cell.number_format = "0.0000000"

    # 自动列宽
    for col_idx in range(1, len(TEMPLATE_HEADERS) + 1):
        max_len = len(TEMPLATE_HEADERS[col_idx - 1])
        for row_idx in range(3, len(rows) + 3):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    wb.save(str(output_path))


# ---------------------------------------------------------------------------
# 日期解析
# ---------------------------------------------------------------------------

def _date_from_yyyymmdd(date_str: str) -> dt.date:
    """将 YYYYMMDD 字符串转为 dt.date。"""
    return dt.datetime.strptime(date_str, "%Y%m%d").date()


def _short_yymmdd(date_str: str) -> str:
    """将 YYYYMMDD 转为 YYMMDD。"""
    return date_str[2:]


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def process_rate_export(
    date_str: str | None = None,
    mailbox: str | None = None,
    output_dir: str | None = None,
    send_error_notify: bool = True,
    include_hkex: bool = True,
    capture_screenshots: bool = True,
) -> dict:
    """从 BOCHK 获取汇率并按 TFISF Excel 公式计算自定义汇率，生成 xlsx 文件。

    Args:
        date_str: YYYYMMDD 格式日期，None 表示今天
        mailbox: 保留兼容，当前未使用
        output_dir: 输出目录，默认 ./temp
        send_error_notify: 出错时是否发送企业微信通知，默认 True
        include_hkex: 是否抓取 HKEx 印花税率并生成交易所汇率.xlsx，
            默认 True（完整方案）；False 时只抓 BOCHK，只生成自定义汇率
        capture_screenshots: 是否截图 BOCHK 牌价页面作为邮件附件，
            默认 True；截图失败只记录警告，不影响主流程

    Returns:
        包含以下键的字典：
          - date_str: 日期字符串 YYYYMMDD
          - date_val: dt.date 日期对象
          - custom_path: 自定义汇率 xlsx 文件路径
          - exchange_path: 交易所汇率 xlsx 文件路径
          - custom_rows: 自定义汇率行列表
          - exchange_rows: 交易所汇率行列表
          - bochk_hkd_raw: BOCHK 港币原始汇率字典
          - bochk_usd_raw: BOCHK 美元原始汇率字典
          - hkex_data: HKEx 原始数据字典
          - screenshots: BOCHK 页面截图列表，每项含 label / url / path
          - errors: 错误信息列表
    """
    if date_str is None:
        date_str = dt.date.today().strftime("%Y%m%d")

    date_val = _date_from_yyyymmdd(date_str)
    yymmdd = _short_yymmdd(date_str)

    out_dir = Path(output_dir) if output_dir else Path.cwd() / "temp"
    out_dir.mkdir(parents=True, exist_ok=True)

    logger.info("[RE] 开始汇率整合导出, date=%s", date_str)

    errors: list[str] = []

    bochk_rates = _collect_bochk_rates(date_str, errors)
    bochk_hkd_raw = _bochk_to_raw_dict(bochk_rates)
    logger.info("[RE] BOCHK-HKD: %d 种货币", len(bochk_hkd_raw))

    bochk_usd_rates = _collect_bochk_usd_rates(date_str, errors)
    bochk_usd_raw = _bochk_usd_to_raw_dict(bochk_usd_rates)
    logger.info("[RE] BOCHK-USD: %d 个货币对", len(bochk_usd_raw))

    bochk_fx_rates = _collect_bochk_fx_rates(date_str, errors)
    bochk_fx_raw = _bochk_fx_to_raw_dict(bochk_fx_rates)
    logger.info("[RE] BOCHK-FX: %d 种货币", len(bochk_fx_raw))

    hkex_data: dict = {}
    hkex_rows: list[dict] = []
    if include_hkex:
        hkex_data = _collect_hkex_rates(date_str, errors)
        hkex_rows = _hkex_to_rows(hkex_data)
        logger.info("[RE] HKEx: %d 条汇率", len(hkex_rows))
    else:
        logger.info("[RE] 仅 BOCHK 模式: 跳过 HKEx 印花税率抓取")

    custom_rows = _compute_custom_rows(
        bochk_hkd_raw, bochk_usd_raw, bochk_fx_raw,
    )
    custom_rows = _filter_rows(custom_rows)
    custom_rows = _sort_rows(custom_rows)

    missing = _missing_pairs(custom_rows)
    if missing:
        msg = (
            "RATE_PAIRS 中配置的币对未取到数据: " + ", ".join(missing) +
            "（BOCHK 电汇/美元/现钞牌价均无该币对）"
        )
        errors.append(msg)
        logger.warning("[RE] %s", msg)

    exchange_rows = _filter_rows(hkex_rows)
    exchange_rows = _sort_rows(exchange_rows)

    custom_path = out_dir / f"自定义汇率{yymmdd}.xlsx"
    _write_xlsx(custom_rows, custom_path, date_val)
    logger.info("[RE] 生成自定义汇率: %s (%d 行)", custom_path, len(custom_rows))
    print(f"[RE] 生成自定义汇率: {custom_path} ({len(custom_rows)} 行)")

    # 仅在有实际交易所汇率时生成交易所汇率.xlsx（非交易日/HKEx 无数据时
    # 不生成该文件，也不作为附件发送，避免误导）。
    exchange_path = None
    if exchange_rows:
        exchange_path = out_dir / f"交易所汇率{yymmdd}.xlsx"
        _write_xlsx(exchange_rows, exchange_path, date_val)
        logger.info(
            "[RE] 生成交易所汇率: %s (%d 行)", exchange_path, len(exchange_rows)
        )
        print(f"[RE] 生成交易所汇率: {exchange_path} ({len(exchange_rows)} 行)")
    else:
        logger.info(
            "[RE] 无交易所汇率数据(HKEx 无数据/非交易日), 不生成交易所汇率.xlsx"
        )

    # BOCHK 页面截图（作为邮件附件，供人工核对当时页面数字）
    # 截图失败只告警，不阻断邮件发送
    screenshots: list[dict] = []
    if capture_screenshots:
        screenshots = capture_rate_pages(date_str, out_dir)
        logger.info("[RE] BOCHK 页面截图: %d 张", len(screenshots))
        print(f"[RE] BOCHK 页面截图: {len(screenshots)} 张")
    else:
        logger.info("[RE] 已禁用页面截图")

    if errors and send_error_notify:
        error_lines = [f"【汇率导出异常 {date_str}】"]
        for e in errors:
            error_lines.append(f"  - {e}")
        try:
            send_error("\n".join(error_lines))
        except Exception as exc:
            logger.error("[RE] 企业微信错误通知发送失败: %s", exc)

    return {
        "date_str": date_str,
        "date_val": date_val,
        "custom_path": custom_path,
        "exchange_path": exchange_path,
        "custom_rows": custom_rows,
        "exchange_rows": exchange_rows,
        "bochk_hkd_raw": bochk_hkd_raw,
        "bochk_usd_raw": bochk_usd_raw,
        "bochk_fx_raw": bochk_fx_raw,
        "hkex_data": hkex_data,
        "screenshots": screenshots,
        "errors": errors,
    }
